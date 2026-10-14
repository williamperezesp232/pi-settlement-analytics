"""
02_analyze.py
Analytics pipeline: CSV -> business metrics (mirror of sql/02_metrics_view.sql) -> executive Excel + charts.

Outputs
  data/pi_cases_enriched.csv            case-level metrics (feed for Power BI)
  reports/PI_Executive_Report.xlsx      KPI workbook (summary tabs are live formulas over the CaseData tab)
  reports/charts/*.png                  figures for the README / Notion

Run:  python python/02_analyze.py
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CSV_IN = ROOT / "data" / "pi_cases.csv"
CSV_OUT = ROOT / "data" / "pi_cases_enriched.csv"
XLSX_OUT = ROOT / "reports" / "PI_Executive_Report.xlsx"
CHART_DIR = ROOT / "reports" / "charts"

AS_OF_DATE = pd.Timestamp("2026-10-13")   
STAGNANT_THRESHOLD_DAYS = 45
DATE_COLS = ["intake_date", "treatment_completion_date", "medical_records_received_date",
             "demand_sent_date", "first_offer_date", "settlement_accepted_date", "closing_statement_date"]



def days(later: pd.Series, earlier: pd.Series) -> pd.Series:
    return (later - earlier).dt.days.astype("Int64")


def add_metrics(df: pd.DataFrame, as_of: pd.Timestamp = AS_OF_DATE) -> pd.DataFrame:
    d = df.copy()

    # Phase durations (NULL until the milestone is reached)
    d["treatment_days"] = days(d.treatment_completion_date, d.intake_date)                 # client / medical driven
    d["records_wait_days"] = days(d.medical_records_received_date, d.treatment_completion_date)  # EXTERNAL
    d["demand_prep_days"] = days(d.demand_sent_date, d.medical_records_received_date)      # CONTROLLABLE
    d["offer_wait_days"] = days(d.first_offer_date, d.demand_sent_date)                    # EXTERNAL
    d["negotiation_days"] = days(d.settlement_accepted_date, d.first_offer_date)           # strategic / shared
    d["closing_days"] = days(d.closing_statement_date, d.settlement_accepted_date)         # CONTROLLABLE
    d["intake_to_demand_days"] = days(d.demand_sent_date, d.intake_date)                   # gross (incl. treatment + records)
    d["total_cycle_days"] = days(d.closing_statement_date, d.intake_date)

    # Controllable vs External
    d["controllable_days"] = d.demand_prep_days + d.closing_days
    d["external_wait_days"] = d.records_wait_days + d.offer_wait_days
    d["controllable_pct_of_cycle"] = (100 * d.controllable_days / d.total_cycle_days).astype(float).round(1)
    d["external_pct_of_cycle"] = (100 * d.external_wait_days / d.total_cycle_days).astype(float).round(1)

    # Ongoing waits on open cases
    rec_open = d.treatment_completion_date.notna() & d.medical_records_received_date.isna()
    off_open = d.demand_sent_date.notna() & d.first_offer_date.isna()
    d["records_wait_open_days"] = days(pd.Series(as_of, index=d.index), d.treatment_completion_date).where(rec_open)
    d["offer_wait_open_days"] = days(pd.Series(as_of, index=d.index), d.demand_sent_date).where(off_open)

    
    d["case_status"] = np.where(d.closing_statement_date.notna(), "Closed", "Open")
    d["current_stage"] = np.select(
        [d.closing_statement_date.notna(), d.settlement_accepted_date.notna(), d.first_offer_date.notna(),
         d.demand_sent_date.notna(), d.medical_records_received_date.notna(), d.treatment_completion_date.notna()],
        ["7. Closed", "6. Closing / Disbursement", "5. Negotiation", "4. Awaiting Adjuster Offer",
         "3. Demand Preparation", "2. Awaiting Records/Bills"],
        default="1. In Treatment")

    
    d["lien_reduction_savings"] = (d.original_liens_amount - d.final_reduced_liens_amount).round(2)
    d["lien_reduction_pct"] = (100 * d.lien_reduction_savings / d.original_liens_amount.replace(0, np.nan)).round(1)

    
    d["net_attorney_fee"] = (d.settlement_amount * d.attorney_fee_pct).round(2)
    d["client_net_recovery_before_costs"] = (d.settlement_amount - d.net_attorney_fee - d.final_reduced_liens_amount).round(2)

    
    rec_eff = d.records_wait_days.fillna(d.records_wait_open_days)
    off_eff = d.offer_wait_days.fillna(d.offer_wait_open_days)
    rec_flag = (rec_eff > STAGNANT_THRESHOLD_DAYS).fillna(False).astype(bool)
    off_flag = (off_eff > STAGNANT_THRESHOLD_DAYS).fillna(False).astype(bool)
    d["stagnant_case_flag"] = rec_flag | off_flag                                  # historical OR active
    d["active_stagnation_alert"] = ((d.records_wait_open_days > STAGNANT_THRESHOLD_DAYS).fillna(False)
                                    | (d.offer_wait_open_days > STAGNANT_THRESHOLD_DAYS).fillna(False)).astype(bool)
    d["stagnation_type"] = np.select([rec_flag & off_flag, rec_flag, off_flag],
                                     ["Records + Adjuster", "Records/Bills", "Adjuster Offer"], default="")
    return d



def console_summary(d: pd.DataFrame) -> None:
    closed = d[d.case_status == "Closed"]
    print("=" * 70, "\nPI SETTLEMENT & BOTTLENECK ANALYTICS  |  as of", AS_OF_DATE.date())
    print("=" * 70)
    print(f"Cases: {len(d)} | Closed: {len(closed)} | Open: {(d.case_status == 'Open').sum()} "
          f"| ACTIVE stagnation alerts: {int(d.active_stagnation_alert.sum())}")
    print(f"Avg cycle (closed): {closed.total_cycle_days.mean():.0f} d | median {closed.total_cycle_days.median():.0f} d")
    print(f"Controllable share of cycle: {100 * closed.controllable_days.sum() / closed.total_cycle_days.sum():.1f}% | "
          f"External share: {100 * closed.external_wait_days.sum() / closed.total_cycle_days.sum():.1f}%")
    print(f"Lien reduction (closed, $-weighted): {100 * closed.lien_reduction_savings.sum() / closed.original_liens_amount.sum():.1f}% "
          f"= ${closed.lien_reduction_savings.sum():,.0f} saved for clients")
    print(f"Attorney fees (all accepted settlements): ${d.net_attorney_fee.sum():,.0f}")
    print(f"Historical stagnation rate: {100 * d.stagnant_case_flag.mean():.1f}% of cases")
    print("\nActive alerts:")
    cols = ["case_id", "case_manager", "insurance_company", "current_stage", "stagnation_type"]
    alerts = d[d.active_stagnation_alert].assign(
        days_without_progress=lambda x: x.records_wait_open_days.fillna(x.offer_wait_open_days))
    print(alerts[cols + ["days_without_progress"]].sort_values("days_without_progress", ascending=False).to_string(index=False))



FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
TITLE = Font(name=FONT, bold=True, size=14, color="1F3864")
NOTE = Font(name=FONT, italic=True, size=9, color="595959")
ALERT_FILL = PatternFill("solid", start_color="FCE4D6")

DATA_COLS = [
    "case_id", "client_name", "injury_type", "case_manager", "insurance_company", "case_status", "current_stage",
    "total_medical_bills", "settlement_amount", "attorney_fee_pct", "original_liens_amount",
    "final_reduced_liens_amount", "intake_date", "treatment_completion_date", "medical_records_received_date",
    "demand_sent_date", "first_offer_date", "settlement_accepted_date", "closing_statement_date",
    "treatment_days", "records_wait_days", "demand_prep_days", "offer_wait_days", "negotiation_days",
    "closing_days", "intake_to_demand_days", "total_cycle_days", "controllable_days", "external_wait_days",
    "lien_reduction_savings", "lien_reduction_pct", "net_attorney_fee", "client_net_recovery_before_costs",
    "stagnant_case_flag", "active_stagnation_alert", "stagnation_type",
]
MONEY = {"total_medical_bills", "settlement_amount", "original_liens_amount", "final_reduced_liens_amount",
         "lien_reduction_savings", "net_attorney_fee", "client_net_recovery_before_costs"}


def style_header(ws, row: int, ncols: int, start_col: int = 1) -> None:
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_excel(d: pd.DataFrame) -> None:
    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
    data = d[DATA_COLS].copy()
    for c in ("stagnant_case_flag", "active_stagnation_alert"):
        data[c] = np.where(data[c], "Yes", "No")
    n = len(data)

    alerts = d[d.active_stagnation_alert].assign(
        days_without_progress=lambda x: x.records_wait_open_days.fillna(x.offer_wait_open_days)
    ).sort_values("days_without_progress", ascending=False)[
        ["case_id", "client_name", "case_manager", "insurance_company", "current_stage",
         "stagnation_type", "days_without_progress"]]

    with pd.ExcelWriter(XLSX_OUT, engine="openpyxl", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as xw:
        pd.DataFrame().to_excel(xw, sheet_name="Executive Summary")
        pd.DataFrame().to_excel(xw, sheet_name="Phase Breakdown")
        pd.DataFrame().to_excel(xw, sheet_name="Case Manager")
        pd.DataFrame().to_excel(xw, sheet_name="Carrier")
        pd.DataFrame().to_excel(xw, sheet_name="Injury Type")
        alerts.to_excel(xw, sheet_name="Active Alerts", index=False)
        data.to_excel(xw, sheet_name="CaseData", index=False)

    wb = load_workbook(XLSX_OUT)
    col = {name: get_column_letter(i + 1) for i, name in enumerate(DATA_COLS)}

    def R(name: str) -> str:                       # absolute range of a CaseData column
        return f"CaseData!${col[name]}$2:${col[name]}${n + 1}"

    
    ws = wb["CaseData"]
    style_header(ws, 1, len(DATA_COLS))
    for i, name in enumerate(DATA_COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 20 if name in ("client_name", "current_stage") else 16
        for r in range(2, n + 2):
            cell = ws.cell(row=r, column=i)
            cell.font = BODY
            if name in MONEY:
                cell.number_format = '$#,##0.00'
            elif name == "attorney_fee_pct":
                cell.number_format = '0.00'
            elif name == "lien_reduction_pct":
                cell.number_format = '0.0'
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DATA_COLS))}{n + 1}"

    
    ws = wb["Active Alerts"]
    style_header(ws, 1, 7)
    for r in range(2, ws.max_row + 1):
        for c in range(1, 8):
            ws.cell(row=r, column=c).font = BODY
            ws.cell(row=r, column=c).fill = ALERT_FILL
    for i, w in enumerate([14, 22, 20, 18, 28, 20, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.cell(row=ws.max_row + 2, column=1,
            value=f"Open cases with > {STAGNANT_THRESHOLD_DAYS} days waiting on records/bills or an adjuster's first offer, as of {AS_OF_DATE.date()}.").font = NOTE

    
    ws = wb["Executive Summary"]
    ws["A1"], ws["A1"].font = "Personal Injury Settlement & Bottleneck Analytics", TITLE
    ws["A2"], ws["A2"].font = f"Synthetic data | As of {AS_OF_DATE.date()} | Summary values are live formulas over the CaseData tab", NOTE
    ws.append([])
    ws.append(["KPI", "Value", "Definition"])
    style_header(ws, 4, 3)
    kpis = [
        ("Total cases", f"=COUNTA({R('case_id')})", "#,##0", "All intakes in dataset"),
        ("Closed cases", f'=COUNTIF({R("case_status")},"Closed")', "#,##0", "Closing statement signed"),
        ("Open cases", f'=COUNTIF({R("case_status")},"Open")', "#,##0", "No closing statement yet"),
        ("ACTIVE stagnation alerts", f'=COUNTIF({R("active_stagnation_alert")},"Yes")', "#,##0",
         "Open cases waiting > 45 days on records/bills or adjuster's first offer"),
        ("Historical stagnation rate", f'=COUNTIF({R("stagnant_case_flag")},"Yes")/COUNTA({R("case_id")})', "0.0%",
         "Cases that ever exceeded 45 days waiting on records or first offer"),
        ("Avg cycle, intake to closing (days)", f"=AVERAGE({R('total_cycle_days')})", "0", "Closed cases only"),
        ("Median cycle (days)", f"=MEDIAN({R('total_cycle_days')})", "0", "Closed cases only"),
        ("Controllable share of cycle", f"=SUM({R('controllable_days')})/SUM({R('total_cycle_days')})", "0.0%",
         "Demand prep + settlement-to-closing (firm-owned work)"),
        ("External-wait share of cycle", f'=SUMIFS({R("external_wait_days")},{R("case_status")},"Closed")/SUM({R("total_cycle_days")})', "0.0%",
         "Records/bills wait + adjuster first-offer wait (third parties)"),
        ("Gross settlements accepted ($)", f"=SUM({R('settlement_amount')})", "$#,##0", "Includes accepted-not-yet-closed"),
        ("Attorney fees earned ($)", f"=SUM({R('net_attorney_fee')})", "$#,##0", "Settlement x fee % (33% pre-suit / 40% litigated)"),
        ("Original liens, closed cases ($)", f'=SUMIFS({R("original_liens_amount")},{R("case_status")},"Closed")', "$#,##0", "As asserted by lienholders"),
        ("Lien reduction savings ($)", f"=SUM({R('lien_reduction_savings')})", "$#,##0", "Original minus final negotiated liens"),
        ("Lien reduction (dollar-weighted)", "=B17/B16", "0.0%", "Savings / original liens"),
    ]
    for label, formula, fmt, note in kpis:
        ws.append([label, formula, note])
        ws.cell(row=ws.max_row, column=2).number_format = fmt
    for r in range(5, ws.max_row + 1):
        for c in range(1, 4):
            ws.cell(row=r, column=c).font = BODY
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 38, 18, 70
    ws.cell(row=ws.max_row + 2, column=1,
            value="Net_Attorney_Fee = settlement x fee %; case costs are treated as pass-through (not in dataset).").font = NOTE

    
    ws = wb["Phase Breakdown"]
    ws["A1"], ws["A1"].font = "Where does the time go? (average days per phase, closed cases)", TITLE
    ws.append([])
    ws.append(["Phase", "Owner", "Avg days", "% of cycle"])
    style_header(ws, 3, 4)
    phases = [("Treatment (intake to MMI)", "Client / medical", "treatment_days"),
              ("Records & bills wait", "EXTERNAL - providers", "records_wait_days"),
              ("Demand preparation", "CONTROLLABLE - firm", "demand_prep_days"),
              ("Adjuster first-offer wait", "EXTERNAL - carrier", "offer_wait_days"),
              ("Negotiation", "Strategic / shared", "negotiation_days"),
              ("Settlement to closing", "CONTROLLABLE - firm", "closing_days")]
    for label, owner, c in phases:
        ws.append([label, owner, f'=AVERAGEIFS({R(c)},{R("case_status")},"Closed")', None])
    first, last = 4, 3 + len(phases)
    for r in range(first, last + 1):
        ws.cell(row=r, column=4, value=f"=C{r}/SUM($C${first}:$C${last})").number_format = "0.0%"
        ws.cell(row=r, column=3).number_format = "0.0"
    ws.append(["Total", "", f"=SUM(C{first}:C{last})", f"=SUM(D{first}:D{last})"])
    ws.cell(row=last + 1, column=3).number_format, ws.cell(row=last + 1, column=4).number_format = "0.0", "0.0%"
    for r in range(3, last + 2):
        for c in range(1, 5):
            if r > 3:
                ws.cell(row=r, column=c).font = Font(name=FONT, size=10, bold=(r == last + 1))
    for letter, w in zip("ABCD", (32, 26, 12, 12)):
        ws.column_dimensions[letter].width = w
    ch = BarChart()
    ch.type, ch.title, ch.style = "bar", "Avg days per phase (closed cases)", 10
    ch.add_data(Reference(ws, min_col=3, min_row=3, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    ch.legend, ch.height, ch.width = None, 8, 18
    ws.add_chart(ch, "F3")

    
    ws = wb["Case Manager"]
    ws["A1"], ws["A1"].font = "Case manager scorecard (firm-controllable efficiency + lien results)", TITLE
    ws.append([])
    ws.append(["Case manager", "Cases", "Avg demand prep (d)", "Avg settlement-to-closing (d)", "Avg controllable days",
               "Lien reduction % ($-weighted)", "Lien savings ($)", "Attorney fees ($)"])
    style_header(ws, 3, 8)
    for i, m in enumerate(sorted(d.case_manager.unique()), start=4):
        cm = R("case_manager")
        ws.append([m, f'=COUNTIF({cm},A{i})',
                   f"=IFERROR(AVERAGEIFS({R('demand_prep_days')},{cm},A{i}),\"\")",
                   f"=IFERROR(AVERAGEIFS({R('closing_days')},{cm},A{i}),\"\")",
                   f"=IFERROR(AVERAGEIFS({R('controllable_days')},{cm},A{i}),\"\")",
                   f'=IFERROR(SUMIFS({R("lien_reduction_savings")},{cm},A{i})/SUMIFS({R("original_liens_amount")},{cm},A{i},{R("case_status")},"Closed"),"")',
                   f"=SUMIFS({R('lien_reduction_savings')},{cm},A{i})",
                   f"=SUMIFS({R('net_attorney_fee')},{cm},A{i})"])
        for c, f in zip(range(3, 9), ("0.0", "0.0", "0.0", "0.0%", "$#,##0", "$#,##0")):
            ws.cell(row=i, column=c).number_format = f
    for r in range(4, ws.max_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).font = BODY
    for i, w in enumerate([20, 8, 18, 22, 18, 22, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    
    ws = wb["Carrier"]
    ws["A1"], ws["A1"].font = "Carrier response time (external wait) - ILLUSTRATIVE synthetic data", TITLE
    ws.append([])
    ws.append(["Insurance company", "Cases w/ first offer", "Avg days demand to offer", "% offers > 45 days", "Avg settlement ($)"])
    style_header(ws, 3, 5)
    for i, ccar in enumerate(sorted(d.insurance_company.unique()), start=4):
        ic, ow = R("insurance_company"), R("offer_wait_days")
        ws.append([ccar, f'=COUNTIFS({ic},A{i},{ow},">=0")',
                   f"=IFERROR(AVERAGEIFS({ow},{ic},A{i}),\"\")",
                   f'=IFERROR(COUNTIFS({ic},A{i},{ow},">45")/B{i},"")',
                   f"=IFERROR(AVERAGEIFS({R('settlement_amount')},{ic},A{i}),\"\")"])
        for c, f in zip(range(3, 6), ("0.0", "0.0%", "$#,##0")):
            ws.cell(row=i, column=c).number_format = f
    for r in range(4, ws.max_row + 1):
        for c in range(1, 6):
            ws.cell(row=r, column=c).font = BODY
    for i, w in enumerate([22, 20, 24, 18, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    
    ws = wb["Injury Type"]
    ws["A1"], ws["A1"].font = "Injury type economics", TITLE
    ws.append([])
    ws.append(["Injury type", "Cases", "Avg billed ($)", "Avg settlement ($)", "Avg cycle (days)", "% litigated (40% fee)"])
    style_header(ws, 3, 6)
    for i, t in enumerate(["Motor Vehicle Accident", "Slip & Fall", "Pedestrian", "Dog Bite"], start=4):
        it = R("injury_type")
        ws.append([t, f"=COUNTIF({it},A{i})",
                   f"=IFERROR(AVERAGEIFS({R('total_medical_bills')},{it},A{i}),\"\")",
                   f"=IFERROR(AVERAGEIFS({R('settlement_amount')},{it},A{i}),\"\")",
                   f"=IFERROR(AVERAGEIFS({R('total_cycle_days')},{it},A{i}),\"\")",
                   f"=COUNTIFS({it},A{i},{R('attorney_fee_pct')},0.4)/B{i}"])
        for c, f in zip(range(3, 7), ("$#,##0", "$#,##0", "0", "0.0%")):
            ws.cell(row=i, column=c).number_format = f
    for r in range(4, ws.max_row + 1):
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = BODY
    for i, w in enumerate([26, 8, 16, 18, 16, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(XLSX_OUT)



def build_charts(d: pd.DataFrame) -> None:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    closed = d[d.case_status == "Closed"]
    navy, orange, grey = "#1F3864", "#C55A11", "#A6A6A6"

    
    labels = ["Treatment", "Records wait\n(external)", "Demand prep\n(controllable)",
              "Adjuster wait\n(external)", "Negotiation", "Closing\n(controllable)"]
    cols = ["treatment_days", "records_wait_days", "demand_prep_days", "offer_wait_days", "negotiation_days", "closing_days"]
    vals = [closed[c].astype(float).mean() for c in cols]
    colors = [grey, orange, navy, orange, grey, navy]
    fig, ax = plt.subplots(figsize=(9, 4.5))
    ax.bar(labels, vals, color=colors)
    for i, v in enumerate(vals):
        ax.text(i, v + 1, f"{v:.0f}", ha="center", fontsize=10)
    ax.set_ylabel("Avg days (closed cases)")
    ax.set_title("Where the time goes: controllable (navy) vs. external (orange) vs. other (grey)")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout(); fig.savefig(CHART_DIR / "01_phase_breakdown.png", dpi=150); plt.close(fig)

    
    g = closed.groupby("case_manager")[["demand_prep_days", "closing_days"]].mean().astype(float)
    g = g.assign(total=g.sum(axis=1)).sort_values("total")
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.barh(g.index, g.demand_prep_days, color=navy, label="Demand prep")
    ax.barh(g.index, g.closing_days, left=g.demand_prep_days, color="#8FAADC", label="Settlement to closing")
    ax.set_xlabel("Avg controllable days"); ax.set_title("Case manager scorecard: controllable days")
    ax.legend(frameon=False); ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout(); fig.savefig(CHART_DIR / "02_case_manager_controllable.png", dpi=150); plt.close(fig)

    
    c = d.groupby("insurance_company").offer_wait_days.agg(lambda s: s.dropna().astype(float).mean()).sort_values()
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.barh(c.index, c.values, color=orange)
    ax.axvline(STAGNANT_THRESHOLD_DAYS, color="red", ls="--", lw=1); ax.text(STAGNANT_THRESHOLD_DAYS + 0.5, -0.4, "45-day alert", color="red", fontsize=9)
    ax.set_xlabel("Avg days: demand sent to first offer"); ax.set_title("Adjuster response time by carrier (synthetic data)")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout(); fig.savefig(CHART_DIR / "03_carrier_offer_wait.png", dpi=150); plt.close(fig)


def main() -> None:
    df = pd.read_csv(CSV_IN, parse_dates=DATE_COLS)
    d = add_metrics(df)
    d.to_csv(CSV_OUT, index=False, date_format="%Y-%m-%d")
    console_summary(d)
    build_excel(d)
    build_charts(d)
    print(f"\nWrote {CSV_OUT.name}, {XLSX_OUT.name} and charts in {CHART_DIR}")


if __name__ == "__main__":
    main()
